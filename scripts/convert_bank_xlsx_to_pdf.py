from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib.pagesizes import A2, landscape
from reportlab.pdfgen import canvas


SOURCE = Path(r"C:\Users\daudm\OneDrive\Desktop\bank-additional-full.xlsx")
OUTPUT = Path(r"E:\workspaceINTELLIJ\fileprocessApp\outputs\bank-additional-full.pdf")

PAGE_SIZE = landscape(A2)
LEFT_MARGIN = 24
TOP_MARGIN = 24
BOTTOM_MARGIN = 24
FONT_NAME = "Courier"
FONT_SIZE = 5
LINE_HEIGHT = 7


def value_to_text(value):
    if value is None:
        return ""
    return str(value).replace("|", " ").replace("\n", " ").replace("\r", " ").strip()


def draw_header(pdf, page_number):
    width, height = PAGE_SIZE
    pdf.setFont(FONT_NAME, 8)
    pdf.drawString(LEFT_MARGIN, height - TOP_MARGIN, f"bank-additional-full.xlsx - page {page_number}")
    pdf.setFont(FONT_NAME, FONT_SIZE)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]

    pdf = canvas.Canvas(str(OUTPUT), pagesize=PAGE_SIZE)
    width, height = PAGE_SIZE

    y = height - TOP_MARGIN - 16
    page_number = 1
    header_line = None

    draw_header(pdf, page_number)

    for index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        values = [value_to_text(value) for value in row]
        line = "|".join(values)

        if index == 1:
            header_line = line

        if y <= BOTTOM_MARGIN:
            pdf.showPage()
            page_number += 1
            draw_header(pdf, page_number)
            y = height - TOP_MARGIN - 16

            if header_line:
                pdf.drawString(LEFT_MARGIN, y, header_line)
                y -= LINE_HEIGHT

        pdf.drawString(LEFT_MARGIN, y, line)
        y -= LINE_HEIGHT

    pdf.save()
    workbook.close()

    print(OUTPUT)


if __name__ == "__main__":
    main()
